"""Conversões a partir de PDF: imagens, SVG, texto, HTML, Markdown, Word, Excel e PowerPoint."""

from __future__ import annotations

import io
import re
from pathlib import Path
from typing import List, Sequence

import pymupdf
from PIL import Image

from ...domain.job import JobContext, ToolResult
from ...utils.exceptions import DependencyError, PDFError
from .common import opened, parse_pages

IMAGE_FORMATS = {"png": "png", "jpg": "jpeg", "webp": "webp", "tiff": "tiff", "bmp": "bmp"}
_PAGE_BREAK = "\f"


def to_images(job: JobContext, files: Sequence[Path], fmt: str = "png", dpi: int = 150, pages: str = "", gray: bool = False, quality: int = 90, single_tiff: bool = False) -> ToolResult:
    """Exporta páginas como imagens (PNG, JPG, WEBP, TIFF, BMP); TIFF pode ser um único arquivo multipágina."""
    outputs = []
    for index, path in enumerate(files):
        with opened(path) as doc:
            selection = parse_pages(pages, doc.page_count)
            folder = None if (fmt == "tiff" and single_tiff) else job.output_folder(path, "imagens")
            frames: List[Image.Image] = []
            width = len(str(doc.page_count))
            for position, number in enumerate(selection):
                job.progress(index + position / len(selection), len(files), f"{path.name}: página {number + 1}")
                image = _page_image(doc[number], int(dpi), gray)
                if folder is None:
                    frames.append(image)
                    continue
                target = folder / f"{path.stem}_p{str(number + 1).zfill(width)}.{fmt}"
                image.save(target, IMAGE_FORMATS[fmt], **_save_options(fmt, int(quality), int(dpi)))
                outputs.append(target)
            if frames:
                target = job.output_file(path, ext=".tiff")
                frames[0].save(target, "TIFF", save_all=True, append_images=frames[1:], compression="tiff_deflate", dpi=(int(dpi), int(dpi)))
                outputs.append(target)
    return ToolResult(outputs, message=f"{len(outputs)} imagem(ns) gerada(s).")


def _page_image(page: pymupdf.Page, dpi: int, gray: bool) -> Image.Image:
    pix = page.get_pixmap(dpi=dpi, colorspace=pymupdf.csGRAY if gray else pymupdf.csRGB, annots=True)
    return Image.frombytes("L" if gray else "RGB", (pix.width, pix.height), pix.samples)


def _save_options(fmt: str, quality: int, dpi: int) -> dict:
    if fmt in ("jpg", "webp"):
        return {"quality": quality, "dpi": (dpi, dpi)}
    if fmt == "tiff":
        return {"compression": "tiff_deflate", "dpi": (dpi, dpi)}
    return {"dpi": (dpi, dpi)}


def to_svg(job: JobContext, files: Sequence[Path], pages: str = "", text_as_path: bool = True) -> ToolResult:
    """Exporta páginas como SVG vetorial."""
    outputs = []
    for index, path in enumerate(files):
        job.progress(index, len(files), f"Exportando {path.name}")
        with opened(path) as doc:
            folder = job.output_folder(path, "svg")
            for number in parse_pages(pages, doc.page_count):
                target = folder / f"{path.stem}_p{number + 1}.svg"
                target.write_text(doc[number].get_svg_image(text_as_path=text_as_path), encoding="utf-8")
                outputs.append(target)
    return ToolResult(outputs)


def to_text(job: JobContext, files: Sequence[Path], sort: bool = True, page_breaks: bool = False) -> ToolResult:
    """Extrai o texto (UTF-8). Páginas escaneadas sem OCR não têm texto."""
    outputs, empty = [], []
    for index, path in enumerate(files):
        job.progress(index, len(files), f"Extraindo texto de {path.name}")
        with opened(path) as doc:
            separator = _PAGE_BREAK if page_breaks else "\n"
            text = separator.join(page.get_text("text", sort=sort) for page in doc)
        if not text.strip():
            empty.append(path.name)
        target = job.output_file(path, ext=".txt")
        target.write_text(text, encoding="utf-8")
        outputs.append(target)
    note = f"Sem texto (use OCR antes): {', '.join(empty)}" if empty else ""
    return ToolResult(outputs, text=note)


def to_html(job: JobContext, files: Sequence[Path], layout: str = "flow") -> ToolResult:
    """HTML em fluxo (fácil de editar) ou posicionado (fiel ao layout), com imagens incorporadas."""
    mode = "xhtml" if layout == "flow" else "html"
    outputs = []
    for index, path in enumerate(files):
        job.progress(index, len(files), f"Convertendo {path.name}")
        with opened(path) as doc:
            pages = "\n".join(f"<section class='page'>{page.get_text(mode)}</section>" for page in doc)
        style = "body{background:#eee;margin:0;padding:16px} .page{background:#fff;margin:0 auto 16px;padding:24px;max-width:900px;box-shadow:0 1px 4px #0003}"
        target = job.output_file(path, ext=".html")
        target.write_text(f"<!doctype html><html><head><meta charset='utf-8'><title>{path.stem}</title><style>{style}</style></head><body>{pages}</body></html>", encoding="utf-8")
        outputs.append(target)
    return ToolResult(outputs)


def to_markdown(job: JobContext, files: Sequence[Path]) -> ToolResult:
    """Markdown com títulos, listas e tabelas detectados (pymupdf4llm)."""
    try:
        import pymupdf4llm
    except ImportError as error:
        raise DependencyError("Pacote 'pymupdf4llm' não instalado. Execute: pip install pymupdf4llm", dependency="pymupdf4llm") from error
    outputs = []
    for index, path in enumerate(files):
        job.progress(index, len(files), f"Convertendo {path.name}")
        with opened(path) as doc:
            markdown = pymupdf4llm.to_markdown(doc, show_progress=False)
        target = job.output_file(path, ext=".md")
        target.write_text(markdown, encoding="utf-8")
        outputs.append(target)
    return ToolResult(outputs)


def to_word(job: JobContext, files: Sequence[Path], pages: str = "") -> ToolResult:
    """DOCX editável preservando parágrafos, tabelas e imagens (pdf2docx)."""
    try:
        from pdf2docx import Converter
    except ImportError as error:
        raise DependencyError("Pacote 'pdf2docx' não instalado. Execute: pip install pdf2docx", dependency="pdf2docx") from error
    outputs = []
    for index, path in enumerate(files):
        job.progress(index, len(files), f"Convertendo {path.name} (pode demorar)")
        with opened(path) as doc:
            selection = parse_pages(pages, doc.page_count) if pages.strip() else None
        target = job.output_file(path, ext=".docx")
        converter = Converter(str(path))
        try:
            converter.convert(str(target), pages=selection)
        except Exception as error:
            raise PDFError(f"Falha ao converter '{path.name}' para Word: {error}") from error
        finally:
            converter.close()
        outputs.append(target)
    return ToolResult(outputs)


def to_excel(job: JobContext, files: Sequence[Path], mode: str = "tables") -> ToolResult:
    """XLSX com as tabelas detectadas (uma aba por tabela) ou com o texto em linhas/colunas."""
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    outputs, lines = [], []
    for index, path in enumerate(files):
        job.progress(index, len(files), f"Convertendo {path.name}")
        workbook = Workbook()
        workbook.remove(workbook.active)
        with opened(path) as doc:
            for page in doc:
                job.check_cancelled()
                if mode == "tables":
                    for number, table in enumerate(page.find_tables().tables, 1):
                        _write_sheet(workbook, f"Pág{page.number + 1}_Tab{number}", _table_rows(page, table))
                else:
                    rows = [re.split(r"\s{2,}", line.strip()) for line in page.get_text("text", sort=True).splitlines() if line.strip()]
                    _write_sheet(workbook, f"Página {page.number + 1}", rows)
        if not workbook.worksheets:
            _write_sheet(workbook, "Sem tabelas", [["Nenhuma tabela detectada. Tente o modo 'Texto em colunas'."]])
            lines.append(f"{path.name}: nenhuma tabela encontrada")
        else:
            lines.append(f"{path.name}: {len(workbook.worksheets)} aba(s)")
        for sheet in workbook.worksheets:
            for column in sheet.columns:
                width = max((len(str(cell.value or "")) for cell in column), default=8)
                sheet.column_dimensions[get_column_letter(column[0].column)].width = min(60, max(8, width + 2))
        target = job.output_file(path, ext=".xlsx")
        workbook.save(target)
        outputs.append(target)
    return ToolResult(outputs, text="\n".join(lines))


def _table_rows(page: pymupdf.Page, table) -> List[List[str]]:
    """Texto de cada célula recortado pela sua caixa (o `extract()` do PyMuPDF 2.0 separa pontuação em outra linha)."""
    return [[" ".join(page.get_text("text", clip=pymupdf.Rect(cell)).split()) if cell else "" for cell in row.cells] for row in table.rows]


def _write_sheet(workbook, title: str, rows: List[list]) -> None:
    sheet = workbook.create_sheet(title[:31])
    for row in rows:
        sheet.append([_cell(value) for value in row])


def _cell(value):
    if value is None:
        return ""
    text = str(value).replace("\n", " ").strip()
    number = re.fullmatch(r"-?\d{1,3}(\.\d{3})*(,\d+)?|-?\d+(,\d+)?", text)
    if number:
        return float(text.replace(".", "").replace(",", "."))
    return text


def to_powerpoint(job: JobContext, files: Sequence[Path], dpi: int = 150) -> ToolResult:
    """PPTX com cada página como slide (imagem em alta resolução)."""
    try:
        from pptx import Presentation
        from pptx.util import Emu
    except ImportError as error:
        raise DependencyError("Pacote 'python-pptx' não instalado. Execute: pip install python-pptx", dependency="python-pptx") from error
    points_to_emu = 12700
    outputs = []
    for index, path in enumerate(files):
        with opened(path) as doc:
            presentation = Presentation()
            first = doc[0].rect
            presentation.slide_width, presentation.slide_height = Emu(int(first.width * points_to_emu)), Emu(int(first.height * points_to_emu))
            for page in doc:
                job.progress(index + page.number / doc.page_count, len(files), f"{path.name}: slide {page.number + 1}")
                slide = presentation.slides.add_slide(presentation.slide_layouts[6])
                image = io.BytesIO(page.get_pixmap(dpi=int(dpi), annots=True).tobytes("png"))
                slide.shapes.add_picture(image, 0, 0, width=presentation.slide_width, height=presentation.slide_height)
        target = job.output_file(path, ext=".pptx")
        presentation.save(target)
        outputs.append(target)
    return ToolResult(outputs)
